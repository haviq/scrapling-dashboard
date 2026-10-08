import json, sqlite3, time, io, csv, re, hashlib, threading, os
from pathlib import Path
from typing import Optional, List
from fastapi import FastAPI, HTTPException, BackgroundTasks, Request
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import httpx
from concurrent.futures import ThreadPoolExecutor

from gmaps import scrape_gmaps_leads

DB = Path('/app/data/data.db') if Path('/app/data').is_dir() else Path(__file__).parent / 'data.db'
TG_BOT = '8980807692:AAE7oobpuegrE0zKgcd5roZiWTsK6tE3qDY'
TG_CHAT = '1868246682'
SCHED_LOCK = threading.Lock()
AUTH_DEFAULT = 'hans-scrap-2026'

app = FastAPI(title='Scrapling Dashboard v3')
app.add_middleware(CORSMiddleware, allow_origins=['*'], allow_methods=['*'], allow_headers=['*'])

def db_init():
    con = sqlite3.connect(DB)
    con.executescript('''
    CREATE TABLE IF NOT EXISTS scrapes(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        url TEXT, mode TEXT, status TEXT, item_count INTEGER,
        data_json TEXT, error TEXT, created_at TEXT DEFAULT (datetime('now','localtime'))
    );
    CREATE TABLE IF NOT EXISTS schedules(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        url TEXT, mode TEXT, selector TEXT, attr TEXT,
        interval_min INTEGER, last_run TEXT, active INTEGER DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
    CREATE TABLE IF NOT EXISTS batches(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        urls_json TEXT, mode TEXT, selector TEXT, attr TEXT,
        prompt TEXT, paginate INTEGER, status TEXT DEFAULT 'running',
        notify INTEGER DEFAULT 1, done INTEGER DEFAULT 0, total INTEGER,
        created_at TEXT DEFAULT (datetime('now','localtime'))
    );
    ''')
    for col, typ in (('batch_id', 'INTEGER'), ('changed', 'INTEGER')):
        try:
            con.execute('ALTER TABLE scrapes ADD COLUMN ' + col + ' ' + typ)
        except Exception:
            pass
    try:
        con.execute('ALTER TABLE schedules ADD COLUMN last_hash TEXT')
    except Exception:
        pass
    con.commit(); con.close()
db_init()

# ---------- Auth ----------
def get_setting(key, default=None):
    con = sqlite3.connect(DB)
    r = con.execute('SELECT value FROM settings WHERE key=?', (key,)).fetchone()
    con.close()
    return r[0] if r and r[0] is not None else default

def set_setting(key, value):
    con = sqlite3.connect(DB)
    con.execute('INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)', (key, value))
    con.commit(); con.close()

AUTH_EXEMPT = {'/', '/api/auth', '/api/demo/extract', '/api/demo/samples'}

# Public demo: fixed safe URLs so visitors can try it without an account.
DEMO_SAMPLES = [
    {'id': 'quotes', 'label': 'Quotes to Scrape', 'url': 'https://quotes.toscrape.com',
     'prompt': 'Extract every quote with its author and tags'},
    {'id': 'books', 'label': 'Books to Scrape', 'url': 'https://books.toscrape.com',
     'prompt': 'Extract each book title, price and availability'},
    {'id': 'hn', 'label': 'Hacker News', 'url': 'https://news.ycombinator.com',
     'prompt': 'Extract the top stories: rank, title, points and comments'},
]
DEMO_URLS = {s['url'] for s in DEMO_SAMPLES}

# simple in-memory rate limit for the public demo
_demo_hits = {}
_demo_lock = threading.Lock()
DEMO_LIMIT = 5          # requests
DEMO_WINDOW = 600       # seconds

@app.middleware('http')
async def auth_middleware(request: Request, call_next):
    p = request.url.path
    if p.startswith('/api/') and p not in AUTH_EXEMPT:
        tok = get_setting('auth_token', AUTH_DEFAULT)
        provided = request.headers.get('x-token') or request.query_params.get('token')
        if provided != tok:
            return JSONResponse({'error': 'unauthorized'}, status_code=401)
    return await call_next(request)

class AuthReq(BaseModel):
    token: str

@app.post('/api/auth')
def api_auth(req: AuthReq):
    tok = get_setting('auth_token', AUTH_DEFAULT)
    if req.token != tok:
        raise HTTPException(401, 'token salah')
    return {'ok': True, 'token': tok}

class AuthSetReq(BaseModel):
    old: str
    new: str

@app.post('/api/auth/change')
def api_auth_change(req: AuthSetReq):
    tok = get_setting('auth_token', AUTH_DEFAULT)
    if req.old != tok:
        raise HTTPException(401, 'token lama salah')
    if len(req.new) < 6:
        raise HTTPException(400, 'token baru minimal 6 karakter')
    set_setting('auth_token', req.new)
    return {'ok': True}

# ---------- Proxy ----------
def get_random_proxy():
    import random
    raw = get_setting('proxies', '') or ''
    lines = [p.strip() for p in raw.split('\n') if p.strip()]
    return random.choice(lines) if lines else None

class ProxyReq(BaseModel):
    proxies: str

@app.get('/api/settings/proxy')
def get_proxy():
    return {'proxies': get_setting('proxies', '') or ''}

@app.post('/api/settings/proxy')
def set_proxy(req: ProxyReq):
    set_setting('proxies', req.proxies)
    return {'ok': True}

# ---------- Scraping engine ----------
def get_fetchers():
    from scrapling.fetchers import Fetcher, StealthyFetcher, DynamicFetcher
    return Fetcher, StealthyFetcher, DynamicFetcher

def do_fetch(url, mode, timeout=60000):
    Fetcher, StealthyFetcher, DynamicFetcher = get_fetchers()
    proxy = get_random_proxy()
    if mode == 'http':
        return Fetcher.get(url, proxy=proxy, timeout=timeout)
    if mode == 'stealth':
        return StealthyFetcher.fetch(url, headless=True, network_idle=True, proxy=proxy, timeout=timeout)
    return DynamicFetcher.fetch(url, headless=True, network_idle=True, proxy=proxy, timeout=timeout)

def basic_extract(page, selector, attr):
    items = []
    if selector:
        for el in page.css(selector):
            v = el.text if attr == 'text' else el.attrib.get(attr, '')
            if v and str(v).strip():
                items.append(str(v).strip())
        if items:
            return items
    for a in page.css('a'):
        t, href = (a.text or '').strip(), a.attrib.get('href', '')
        if t and href:
            items.append({'text': t[:120], 'href': href})
    for img in page.css('img'):
        src = img.attrib.get('src', '')
        if src and 'loading' not in src:
            items.append({'image': src})
    return items

def smart_extract(page):
    out = {'page_title': '', 'meta_desc': '', 'headings': [], 'menu': [], 'tables': [], 'images': [], 'links': []}
    t = (page.css('title') or [None])[0]
    if t: out['page_title'] = t.text
    m = (page.css('meta[name=description]') or [None])[0]
    if m: out['meta_desc'] = m.attrib.get('content', '')
    for h in page.css('h1, h2, h3'):
        txt = h.text.strip()
        if txt: out['headings'].append({'tag': h.tag, 'text': txt[:150]})
    seen = set()
    for a in page.css('nav a, header a, .menu a, #menu a, .navbar a'):
        label, href = (a.text or '').strip(), a.attrib.get('href', '')
        if label and (label, href) not in seen:
            seen.add((label, href)); out['menu'].append({'label': label[:80], 'href': href})
    for tb in page.css('table'):
        rows = []
        for tr in tb.css('tr'):
            cells = [td.text.strip() for td in tr.css('th, td')]
            if any(cells): rows.append(cells)
        if rows: out['tables'].append(rows)
    seen_img = set()
    for img in page.css('img'):
        src = img.attrib.get('src', '')
        if src and src not in seen_img and 'loading' not in src:
            seen_img.add(src); out['images'].append({'src': src, 'alt': img.attrib.get('alt', '')[:100]})
    seen_l = set()
    for a in page.css('a'):
        label, href = (a.text or '').strip(), a.attrib.get('href', '')
        if href and label and len(label) > 20 and (label, href) not in seen_l:
            seen_l.add((label, href)); out['links'].append({'text': label[:150], 'href': href})
    return {k: v for k, v in out.items() if v}

# ---------- AI config (env-overridable) ----------
AI_BASE_URL = os.environ.get('AI_BASE_URL', 'https://api.example.com/v1')
AI_API_KEY = os.environ.get('AI_API_KEY', '')
AI_MODELS = [m.strip() for m in os.environ.get('AI_MODELS', 'gemini-3.8-flash-high,gemini-3.6-flash-high,gemini-3.5-flash-lite').split(',') if m.strip()]
MODELS = AI_MODELS

def ai_extract(html_text, prompt):
    if isinstance(html_text, (bytes, bytearray)):
        html_text = html_text.decode('utf-8', 'ignore')
    sys = ('You are a web scraping data extractor. Extract structured data as JSON array from the HTML content '
           'based on user instructions. Return ONLY a JSON array, no markdown, no explanation.')
    messages = [{'role': 'system', 'content': sys},
                {'role': 'user', 'content': 'Instruction: ' + prompt + '\n\nHTML (truncated):\n' + html_text[:15000]}]
    headers = {'Authorization': 'Bearer ' + AI_API_KEY,
               'Content-Type': 'application/json'}
    last_err = None
    for attempt in range(1, 4):
        try:
            body = {'model': MODELS[attempt % len(MODELS)], 'messages': messages, 'temperature': 0.1}
            r = httpx.post(AI_BASE_URL + '/chat/completions', json=body, headers=headers, timeout=120)
            r.raise_for_status()
            clean_body = re.sub(r'data:\s*\[DONE\]', '', r.text).strip()
            raw = (json.loads(clean_body)['choices'][0]['message'].get('content') or '').strip()
            if raw.startswith('```'):
                raw = raw.split('```')[1].lstrip('json').strip()
            start_idx, end_idx = raw.find('['), raw.rfind(']')
            if start_idx == -1 or end_idx <= start_idx:
                raise ValueError('no JSON array in response')
            segment = raw[start_idx:]
            try:
                return json.loads(segment)
            except json.JSONDecodeError:
                cleaned = re.sub(r',\s*([\]\}])', r'\1', segment)
                try:
                    return json.loads(cleaned)
                except json.JSONDecodeError:
                    dec = json.JSONDecoder()
                    obj, _ = dec.raw_decode(segment)
                    return obj
        except Exception as e:
            last_err = e
            time.sleep(3)
    raise last_err or RuntimeError('AI extraction failed')

# ---------- Deep scrape (level-2) ----------
def deep_scrape(url, mode, selector, attr, prompt, link_selector, max_links):
    from urllib.parse import urljoin
    page = do_fetch(url, mode)
    if prompt:
        base_items = ai_extract(page.body if hasattr(page, 'body') else page.html_content, prompt)
    elif selector:
        base_items = basic_extract(page, selector, attr)
    else:
        base_items = smart_extract(page)
    links = []
    for a in page.css(link_selector or 'a[href]'):
        href = a.attrib.get('href', '')
        if href and not href.startswith(('#', 'javascript:', 'mailto:', 'tel:')):
            links.append(urljoin(url, href))
    seen = set()
    links = [x for x in links if not (x in seen or seen.add(x))]
    try:
        max_links = max(1, min(int(max_links or 50), 200))
    except Exception:
        max_links = 50
    links = links[:max_links]

    def one(u):
        try:
            pg = do_fetch(u, 'http')
            if prompt:
                return ai_extract(pg.body if hasattr(pg, 'body') else pg.html_content, prompt)
            if selector:
                return basic_extract(pg, selector, attr)
            sm = smart_extract(pg)
            return [sm] if sm else []
        except Exception as e:
            return [{'url': u, 'error': str(e)[:120]}]

    detail = []
    if links:
        with ThreadPoolExecutor(max_workers=8) as ex:
            for res in ex.map(one, links):
                if isinstance(res, list):
                    detail.extend(res)
                elif res:
                    detail.append(res)
    return {'page': base_items, 'detail': detail, 'links_followed': len(links)}

# ---------- Social ----------
def social_scrape(url, cookies):
    from playwright.sync_api import sync_playwright
    results = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(user_agent='Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36')
        if not cookies:
            return [{'error': 'Cookies wajib diisi untuk mode Social Crawler (X/IG/FB).'}]
        try:
            context.add_cookies(json.loads(cookies))
        except Exception:
            return [{'error': 'Format Cookies tidak valid. Harus format JSON dari Cookie-Editor.'}]
        page = context.new_page()
        page.goto(url)
        page.wait_for_timeout(5000)
        seen = set()
        for _ in range(5):
            if 'x.com' in url or 'twitter.com' in url:
                for el in page.query_selector_all('[data-testid="tweetText"]'):
                    txt = el.inner_text().replace('\n', ' ')
                    if txt not in seen:
                        seen.add(txt); results.append({'Tweet': txt})
            elif 'instagram.com' in url:
                for el in page.query_selector_all('article img[alt], main img[alt]'):
                    txt = (el.get_attribute('alt') or '').replace('\n', ' ')
                    if txt and txt not in seen and len(txt) > 10:
                        seen.add(txt); results.append({'Caption': txt})
            else:
                for el in page.query_selector_all('p, h1, h2, h3, article, span'):
                    txt = el.inner_text().replace('\n', ' ')
                    if txt and txt not in seen and len(txt) > 20:
                        seen.add(txt); results.append({'Text': txt})
            page.evaluate('window.scrollBy(0, 1000)')
            page.wait_for_timeout(2000)
        browser.close()
    return results

def fetch_and_extract(url, mode, selector=None, attr='text', prompt=None, paginate=1, cookies=None,
                      link_selector=None, max_links=None, deep=False):
    if mode == 'social':
        return social_scrape(url, cookies)
    if mode == 'gmaps':
        keyword = url.split('q=')[-1] if 'q=' in url else url
        return scrape_gmaps_leads(keyword, limit=50)
    if deep:
        return deep_scrape(url, mode, selector, attr, prompt, link_selector, max_links)

    page = do_fetch(url, mode)
    # JSON response detection (BMKG, Epic, dsb.)
    try:
        jd = page.json()
        if jd:
            items_raw = jd if isinstance(jd, list) else [jd]
            flat = []
            def walk(obj):
                if isinstance(obj, dict):
                    flat.append({k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else str(v)) for k, v in obj.items()})
                elif isinstance(obj, list):
                    for x in obj:
                        walk(x)
            walk(items_raw)
            if flat:
                if len(flat) == 1:
                    for v in flat[0].values():
                        try:
                            nested = json.loads(v)
                            if isinstance(nested, list) and nested:
                                flat = []
                                walk(nested)
                                break
                            elif isinstance(nested, dict):
                                for val in nested.values():
                                    if isinstance(val, list) and val and isinstance(val[0], dict):
                                        flat = []
                                        walk(val)
                                        break
                                break
                        except Exception:
                            pass
                return flat
    except Exception:
        pass
    if not selector and not prompt:
        return smart_extract(page)
    if prompt:
        return ai_extract(page.body if hasattr(page, 'body') else page.html_content, prompt)
    items = basic_extract(page, selector, attr)
    try:
        paginate = max(1, min(int(paginate or 1), 50))
    except Exception:
        paginate = 1
    if paginate > 1:
        base = url.rstrip('/')
        patterns = [base + '/{n}', base + '?page={n}', re.sub(r'(-\d+)$', r'-{n}', base)]
        def fetch_page(n):
            for pat in patterns:
                nxt = pat.format(n=n)
                if nxt == url:
                    continue
                try:
                    return basic_extract(do_fetch(nxt, mode), selector, attr)
                except Exception:
                    continue
            return []
        with ThreadPoolExecutor(max_workers=min(10, paginate)) as ex:
            for res in ex.map(fetch_page, range(2, paginate + 1)):
                items.extend(res)
    return items

# ---------- Models ----------
class ScrapeReq(BaseModel):
    url: str
    mode: str = 'http'
    selector: Optional[str] = None
    attr: str = 'text'
    prompt: Optional[str] = None
    cookies: Optional[str] = None
    link_selector: Optional[str] = None
    max_links: int = 50
    deep: bool = False
    paginate: int = 1
    notify: bool = True

class SchedReq(BaseModel):
    url: str
    mode: str = 'http'
    selector: Optional[str] = None
    attr: str = 'text'
    interval_min: int = 60

# ---------- Telegram ----------
def tg_send(text):
    try:
        httpx.post('https://api.telegram.org/bot' + TG_BOT + '/sendMessage',
                   json={'chat_id': TG_CHAT, 'text': text[:4000], 'parse_mode': 'HTML'}, timeout=15)
    except Exception:
        pass

def save_scrape(url, mode, status, items=None, error=None):
    con = sqlite3.connect(DB)
    con.execute('INSERT INTO scrapes(url,mode,status,item_count,data_json,error) VALUES(?,?,?,?,?,?)',
                (url, mode, status, len(items or []), json.dumps(items or [], ensure_ascii=False), (error or '')[:500]))
    sid = con.execute('SELECT last_insert_rowid()').fetchone()[0]
    con.commit(); con.close()
    return sid

def update_scrape(sid, status, items=None, error=None):
    con = sqlite3.connect(DB)
    con.execute('UPDATE scrapes SET status=?, item_count=?, data_json=?, error=? WHERE id=?',
                (status, len(items or []), json.dumps(items or [], ensure_ascii=False), (error or '')[:500], sid))
    con.commit(); con.close()

def job_scrape(sid, url, mode, selector, attr, prompt, paginate, notify, cookies=None,
               link_selector=None, max_links=None, deep=False):
    try:
        items = fetch_and_extract(url, mode, selector, attr, prompt, paginate, cookies,
                                  link_selector, max_links, deep)
        update_scrape(sid, 'done', items)
        if notify:
            tg_send('🕷️ <b>Scrape selesai #' + str(sid) + '</b>\nURL: ' + url + '\nItems: ' + str(len(items)) + '\n\nhttps://scrapdash.haviq.dev')
    except Exception as e:
        update_scrape(sid, 'error', error=str(e))
        if notify:
            tg_send('❌ <b>Scrape gagal #' + str(sid) + '</b>\n' + url + '\n' + str(e)[:200])

# ---------- Routes ----------
@app.post('/api/scrape')
def api_scrape(req: ScrapeReq, bg: BackgroundTasks):
    sid = save_scrape(req.url, req.mode, 'running')
    bg.add_task(job_scrape, sid, req.url, req.mode, req.selector, req.attr, req.prompt,
                req.paginate, req.notify, req.cookies, req.link_selector, req.max_links, req.deep)
    return {'id': sid, 'status': 'running'}

# ---------- Bulk Queue ----------
class BulkReq(BaseModel):
    urls: List[str]
    mode: str = 'http'
    selector: Optional[str] = None
    attr: str = 'text'
    prompt: Optional[str] = None
    cookies: Optional[str] = None
    paginate: int = 1
    notify: bool = True

@app.post('/api/bulk')
def api_bulk(req: BulkReq, bg: BackgroundTasks):
    urls = [u.strip() for u in req.urls if u.strip()][:100]
    if not urls:
        raise HTTPException(400, 'urls kosong')
    con = sqlite3.connect(DB)
    cur = con.execute('INSERT INTO batches(urls_json,mode,selector,attr,prompt,paginate,notify,total) VALUES(?,?,?,?,?,?,?,?)',
                      (json.dumps(urls), req.mode, req.selector, req.attr, req.prompt, req.paginate,
                       1 if req.notify else 0, len(urls)))
    bid = cur.lastrowid
    con.commit(); con.close()
    bg.add_task(run_batch, bid)
    return {'batch_id': bid, 'total': len(urls), 'status': 'running'}

def run_batch(bid):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    b = con.execute('SELECT * FROM batches WHERE id=?', (bid,)).fetchone()
    con.close()
    if not b:
        return
    urls = json.loads(b['urls_json'])
    done, failed = 0, 0
    for u in urls:
        sid = save_scrape(u, b['mode'], 'running')
        con = sqlite3.connect(DB)
        con.execute('UPDATE scrapes SET batch_id=? WHERE id=?', (bid, sid))
        con.commit(); con.close()
        try:
            items = fetch_and_extract(u, b['mode'], b['selector'], b['attr'], b['prompt'], b['paginate'])
            update_scrape(sid, 'done', items)
        except Exception as e:
            update_scrape(sid, 'error', error=str(e))
            failed += 1
        done += 1
        con = sqlite3.connect(DB)
        con.execute('UPDATE batches SET done=? WHERE id=?', (done, bid))
        con.commit(); con.close()
    con = sqlite3.connect(DB)
    con.execute("UPDATE batches SET status='done' WHERE id=?", (bid,))
    con.commit(); con.close()
    if b['notify']:
        tg_send('📦 <b>Batch #' + str(bid) + ' selesai</b>\n' + str(done - failed) + '/' + str(len(urls)) + ' sukses\nhttps://scrapdash.haviq.dev')

@app.get('/api/batches')
def api_batches():
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    rows = con.execute('SELECT id,status,done,total,mode,created_at FROM batches ORDER BY id DESC LIMIT 50').fetchall()
    con.close()
    return [dict(r) for r in rows]

@app.get('/api/batch/{bid}')
def api_batch_detail(bid: int):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    b = con.execute('SELECT id,status,done,total,mode,created_at FROM batches WHERE id=?', (bid,)).fetchone()
    if not b:
        raise HTTPException(404)
    items = con.execute('SELECT id,url,status,item_count FROM scrapes WHERE batch_id=? ORDER BY id', (bid,)).fetchall()
    con.close()
    return {'batch': dict(b), 'scrapes': [dict(x) for x in items]}

# ---------- Analyze (Chat-to-Data) ----------
class AnalyzeReq(BaseModel):
    prompt: str

@app.post('/api/analyze/{sid}')
def analyze_data(sid: int, req: AnalyzeReq):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    r = con.execute('SELECT data_json FROM scrapes WHERE id=?', (sid,)).fetchone(); con.close()
    if not r:
        raise HTTPException(404)
    data_text = (r['data_json'] or '[]')[:15000]
    sys = 'You are a data analyst. Answer the user prompt based on the provided JSON data. Be concise and insightful. Use Markdown.'
    messages = [{'role': 'system', 'content': sys},
                {'role': 'user', 'content': 'Prompt: ' + req.prompt + '\n\nData:\n' + data_text}]
    headers = {'Authorization': 'Bearer ' + AI_API_KEY,
               'Content-Type': 'application/json'}
    try:
        body = {'model': MODELS[0], 'messages': messages, 'temperature': 0.3}
        resp = httpx.post(AI_BASE_URL + '/chat/completions', json=body, headers=headers, timeout=120)
        resp.raise_for_status()
        clean_body = re.sub(r'data:\s*\[DONE\]', '', resp.text).strip()
        raw = (json.loads(clean_body)['choices'][0]['message'].get('content') or '').strip()
        return {'answer': raw}
    except Exception as e:
        return {'error': str(e)}

# ---------- AI Extract (URL -> structured JSON) ----------
class AIExtractReq(BaseModel):
    url: str
    prompt: str
    mode: str = 'http'

@app.post('/api/ai-extract')
def api_ai_extract(req: AIExtractReq):
    try:
        page = do_fetch(req.url, req.mode)
    except Exception as e:
        raise HTTPException(400, 'fetch failed: ' + str(e))
    html = page.body if hasattr(page, 'body') else getattr(page, 'html_content', '')
    try:
        items = ai_extract(html, req.prompt)
    except Exception as e:
        raise HTTPException(502, 'ai extract failed: ' + str(e))
    return {'url': req.url, 'model': MODELS[0], 'count': len(items) if isinstance(items, list) else 0, 'items': items}

# ---------- Public demo (no auth) ----------
@app.get('/api/demo/samples')
def demo_samples():
    return [{'id': s['id'], 'label': s['label'], 'prompt': s['prompt']} for s in DEMO_SAMPLES]

class DemoReq(BaseModel):
    sample: str

@app.post('/api/demo/extract')
def demo_extract(req: DemoReq, request: Request):
    sample = next((s for s in DEMO_SAMPLES if s['id'] == req.sample), None)
    if not sample:
        raise HTTPException(400, 'unknown sample')
    ip = (request.client.host if request.client else 'unknown')
    now = time.time()
    with _demo_lock:
        hits = [t for t in _demo_hits.get(ip, []) if now - t < DEMO_WINDOW]
        if len(hits) >= DEMO_LIMIT:
            raise HTTPException(429, 'demo rate limit reached, try again later')
        hits.append(now)
        _demo_hits[ip] = hits
    try:
        resp = httpx.get(sample['url'], timeout=30, follow_redirects=True,
                         headers={'User-Agent': 'Mozilla/5.0 (compatible; ScraplingDemo/1.0)'})
        resp.raise_for_status()
        html = resp.text
        items = ai_extract(html, sample['prompt'])
    except Exception as e:
        raise HTTPException(502, 'demo failed: ' + str(e))
    return {'sample': sample['id'], 'label': sample['label'], 'count': len(items) if isinstance(items, list) else 0, 'items': (items[:8] if isinstance(items, list) else items)}

# ---------- History ----------
@app.get('/api/history')
def history(q: str = '', status: str = ''):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    sql = 'SELECT id,url,mode,status,item_count,error,created_at,changed FROM scrapes'
    cond, args = [], []
    if q:
        cond.append('url LIKE ?'); args.append('%' + q + '%')
    if status:
        cond.append('status = ?'); args.append(status)
    if cond:
        sql += ' WHERE ' + ' AND '.join(cond)
    sql += ' ORDER BY id DESC LIMIT 200'
    rows = con.execute(sql, args).fetchall()
    con.close()
    return [dict(r) for r in rows]

@app.get('/api/history/{sid}')
def detail(sid: int):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    r = con.execute('SELECT * FROM scrapes WHERE id=?', (sid,)).fetchone(); con.close()
    if not r:
        raise HTTPException(404)
    d = dict(r); d['data'] = json.loads(d['data_json'] or '[]')
    return d

@app.delete('/api/history/{sid}')
def del_scrape(sid: int):
    con = sqlite3.connect(DB)
    con.execute('DELETE FROM scrapes WHERE id=?', (sid,))
    con.commit(); con.close()
    return {'ok': True}

@app.get('/api/stats')
def stats():
    con = sqlite3.connect(DB)
    total = con.execute('SELECT COUNT(*), COALESCE(SUM(item_count),0) FROM scrapes').fetchone()
    done = con.execute("SELECT COUNT(*) FROM scrapes WHERE status='done'").fetchone()[0]
    changed = con.execute('SELECT COUNT(*) FROM scrapes WHERE changed=1').fetchone()[0]
    con.close()
    return {'total': total[0], 'items': total[1], 'done': done, 'changed': changed}

# ---------- Download ----------
@app.get('/api/download/{sid}')
def download(sid: int, format: str = 'csv'):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    r = con.execute('SELECT * FROM scrapes WHERE id=?', (sid,)).fetchone(); con.close()
    if not r:
        raise HTTPException(404)
    data = json.loads(r['data_json'] or '[]')
    if format == 'json':
        buf = json.dumps(data, ensure_ascii=False, indent=2)
        media, fname = 'application/json', 'scrape_' + str(sid) + '.json'
    elif format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook(); ws = wb.active; ws.title = 'Hasil Scrape'
        norm = [row if isinstance(row, dict) else {'value': str(row)} for row in data]
        keys = []
        for row in norm:
            for k in row.keys():
                if k not in keys:
                    keys.append(k)
        ws.append([k.upper() for k in keys])
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        for row in norm:
            ws.append([str(row.get(k, '')) if not isinstance(row.get(k, ''), (int, float)) else row.get(k) for k in keys])
        for i, k in enumerate(keys, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(15, min(50, len(k) + 15))
        ws.freeze_panes = 'A2'
        buf_io = io.BytesIO(); wb.save(buf_io); buf = buf_io.getvalue()
        media, fname = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'scrape_' + str(sid) + '.xlsx'
    else:
        out = io.StringIO()
        norm = [row if isinstance(row, dict) else {'value': str(row)} for row in data]
        keys = []
        for row in norm:
            for k in row.keys():
                if k not in keys:
                    keys.append(k)
        w = csv.DictWriter(out, fieldnames=keys, extrasaction='ignore')
        w.writeheader()
        for row in norm:
            w.writerow(row)
        buf = out.getvalue(); media, fname = 'text/csv', 'scrape_' + str(sid) + '.csv'
    return StreamingResponse(iter([buf] if isinstance(buf, str) else [buf]), media_type=media,
        headers={'Content-Disposition': 'attachment; filename=' + fname})

@app.post('/api/telegram/{sid}')
def send_tg(sid: int):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    r = con.execute('SELECT * FROM scrapes WHERE id=?', (sid,)).fetchone(); con.close()
    if not r:
        raise HTTPException(404)
    data = json.loads(r['data_json'] or '[]')
    preview = '\n'.join(json.dumps(x, ensure_ascii=False)[:100] if isinstance(x, dict) else str(x)[:100] for x in data[:15])
    tg_send('🕷️ <b>Hasil Scrape #' + str(sid) + '</b>\n' + r['url'] + '\n\n<pre>' + preview + '</pre>\n\nTotal: ' + str(r['item_count']) + ' items')
    return {'ok': True}

# ---------- Scheduler with diff ----------
@app.post('/api/schedules')
def add_schedule(req: SchedReq):
    con = sqlite3.connect(DB)
    con.execute('INSERT INTO schedules(url,mode,selector,attr,interval_min) VALUES(?,?,?,?,?)',
                (req.url, req.mode, req.selector, req.attr, req.interval_min))
    con.commit(); con.close()
    return {'ok': True}

@app.get('/api/schedules')
def list_schedules():
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    rows = con.execute('SELECT * FROM schedules ORDER BY id DESC').fetchall(); con.close()
    return [dict(r) for r in rows]

@app.delete('/api/schedules/{sid}')
def del_schedule(sid: int):
    con = sqlite3.connect(DB)
    con.execute('DELETE FROM schedules WHERE id=?', (sid,))
    con.commit(); con.close()
    return {'ok': True}

@app.get('/', response_class=HTMLResponse)
def dashboard():
    return Path('dashboard.html').read_text()

def scheduler_loop():
    import datetime
    while True:
        try:
            con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
            rows = con.execute('SELECT * FROM schedules WHERE active=1').fetchall(); con.close()
            now = datetime.datetime.now().isoformat(timespec='seconds')
            for r in rows:
                last = r['last_run'] or ''
                due = (not last) or (datetime.datetime.fromisoformat(last) + datetime.timedelta(minutes=r['interval_min']) <= datetime.datetime.now())
                if due:
                    with SCHED_LOCK:
                        con = sqlite3.connect(DB)
                        con.execute('UPDATE schedules SET last_run=? WHERE id=?', (now, r['id']))
                        con.commit(); con.close()
                    try:
                        page = do_fetch(r['url'], r['mode'])
                        items = (smart_extract(page) if not r['selector'] else basic_extract(page, r['selector'], r['attr']))
                        sig = hashlib.md5(json.dumps(items, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
                        con2 = sqlite3.connect(DB); con2.row_factory = sqlite3.Row
                        prev = con2.execute('SELECT last_hash FROM schedules WHERE id=?', (r['id'],)).fetchone()
                        con2.close()
                        changed = (not prev) or (prev['last_hash'] != sig)
                        sid2 = save_scrape(r['url'], r['mode'], 'done', items)
                        con2 = sqlite3.connect(DB)
                        con2.execute('UPDATE schedules SET last_hash=? WHERE id=?', (sig, r['id']))
                        con2.execute('UPDATE scrapes SET changed=? WHERE id=?', (1 if changed else 0, sid2))
                        con2.commit(); con2.close()
                        if changed:
                            tg_send('⏰ <b>Auto-scrape #' + str(r['id']) + '</b>\n' + r['url'] + '\nItems: ' + str(len(items)) + ' 🆕 ADA DATA BARU')
                    except Exception as e:
                        save_scrape(r['url'], r['mode'], 'error', error=str(e))
        except Exception:
            pass
        time.sleep(60)

threading.Thread(target=scheduler_loop, daemon=True).start()
